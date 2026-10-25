"""Build AI_Real_Time_Mapping_Midterm.pptx (python-pptx; native shapes and charts).

Every number below is copied from a committed summary named in SOURCES and in
the speaker notes; nothing is computed from GT here. Run from this directory:
    python build_deck.py
"""
from pathlib import Path

import re
from datetime import datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "AI_Real_Time_Mapping_Midterm.pptx"
COMMIT = "242c0cb"  # origin/develop used for this deck

BG = RGBColor(0x0F, 0x14, 0x19)
CARD = RGBColor(0x1B, 0x23, 0x2C)
CARD2 = RGBColor(0x24, 0x2E, 0x39)
TEXT = RGBColor(0xEB, 0xEE, 0xF0)
MUTED = RGBColor(0x9A, 0xA5, 0xB1)
ORANGE = RGBColor(0xFF, 0xA8, 0x26)  # P1 / DA3 / our pipeline
BLUE = RGBColor(0x5A, 0xA9, 0xE6)  # UniDepth
GREY = RGBColor(0x8A, 0x96, 0xA3)  # B0 sparse
ORACLE = RGBColor(0xC9, 0xD3, 0xDE)  # B1 oracle
RED = RGBColor(0xE5, 0x53, 0x4B)  # worse variant
TEAL = RGBColor(0x40, 0xC4, 0xB4)
PEOPLE = {"Назар": ORANGE, "Ярополк": TEAL, "Юліан": BLUE}
FONT = "Calibri"

prs = Presentation()
properties = prs.core_properties
properties.title = "Real-time Mapping with Foundation Models: проміжна презентація"
properties.author = "Назар, Ярополк, Юліан"
properties.last_modified_by = "Юліан"
properties.comments = f"origin/develop {COMMIT}; offline research prototype"
properties.created = properties.modified = datetime(2026, 10, 25)
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide_base(title, number):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    if title:
        box(slide, 0.6, 0.35, 12.1, 0.8, title, 32, bold=True)
    box(slide, 0.6, 7.0, 8.0, 0.35, "Offline research prototype · UCU · Real-time Mapping with Foundation Models",
        12, color=MUTED)
    box(slide, 11.7, 7.0, 1.0, 0.35, f"{number} / 6", 12, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def box(slide, x, y, w, h, text, size, *, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        runs = line if isinstance(line, list) else [(line, {})]
        for value, style in runs:
            run = paragraph.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(style.get("size", size))
            run.font.bold = style.get("bold", bold)
            run.font.color.rgb = style.get("color", color)
        paragraph.space_after = Pt(style.get("after", 0) if runs else 0)
    return shape


def bullets(slide, x, y, w, h, items, size=20, color=TEXT, gap=8):
    shape = box(slide, x, y, w, h, [[(item, {})] for item in items], size, color=color)
    for paragraph in shape.text_frame.paragraphs:
        paragraph.space_after = Pt(gap)
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Inches(0.3)))
        properties.set("indent", str(-Inches(0.3)))
        properties.append(properties.makeelement(qn("a:buChar"), {"char": "•"}))
    return shape


def card(slide, x, y, w, h, fill=CARD, line=None, radius=0.06, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(2)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = radius
    shape.shadow.inherit = False
    return shape


def node(slide, x, y, w, h, title, sub, owner=None, line=None, dashed=False):
    shape = card(slide, x, y, w, h, CARD2, line=line or (PEOPLE[owner] if owner else MUTED))
    if dashed:
        shape.line.dash_style = 4  # dash
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Inches(0.08)
    for i, (value, size, bold, color) in enumerate([(title, 20, True, TEXT), (sub, 16, False, MUTED)]):
        if not value:
            continue
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text, run.font.size, run.font.bold, run.font.name = value, Pt(size), bold, FONT
        run.font.color.rgb = color
    return shape


def arrow(slide, x1, y1, x2, y2, color=MUTED, dashed=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2.25)
    if dashed:
        line.line.dash_style = 4
    tail = line.line._get_or_add_ln()
    end = tail.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                           {"type": "triangle", "w": "med", "len": "med"})
    tail.append(end)
    return line


def style_chart(chart, number_format, title=None, legend=False, max_value=None, label_size=16, axis_format=None):
    chart.font.name = FONT
    chart.font.size = Pt(16)
    chart.font.color.rgb = TEXT
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(18), True, TEXT, FONT
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(16)
        chart.legend.font.color.rgb = TEXT
    value_axis, category_axis = chart.value_axis, chart.category_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor(0x33, 0x3D, 0x48)
    value_axis.format.line.fill.background()
    value_axis.tick_labels.font.size = Pt(14)
    value_axis.tick_labels.font.color.rgb = MUTED
    value_axis.tick_labels.number_format = axis_format or number_format
    value_axis.tick_labels.number_format_is_linked = False
    value_axis.minimum_scale = 0
    if max_value:
        value_axis.maximum_scale = max_value
    category_axis.tick_labels.font.size = Pt(16)
    category_axis.tick_labels.font.color.rgb = TEXT
    category_axis.format.line.color.rgb = MUTED
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format, labels.number_format_is_linked = number_format, False
    labels.font.size, labels.font.bold, labels.font.color.rgb = Pt(label_size), True, TEXT
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    return chart


def color_points(series, colors):
    for point, color in zip(series.points, colors):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color


SCRIPT = (HERE / "MIDTERM_SCRIPT_UA.md").read_text(encoding="utf-8")


def notes(slide, number, sources):
    """Speaker notes = heading + spoken block of MIDTERM_SCRIPT_UA.md (single source) + sources."""
    heading = re.search(rf"^## Слайд {number} — .*$", SCRIPT, re.M).group(0)[3:]
    spoken = re.search(rf"<!-- notes:{number} -->\n(.*?)<!-- /notes -->", SCRIPT, re.S).group(1).strip()
    slide.notes_slide.notes_text_frame.text = f"{heading}\n\n{spoken}\n\n{sources}"


# ---------------------------------------------------------------- slide 1
s = slide_base(None, 1)
box(s, 0.6, 0.45, 5.6, 0.5, "Real-time Mapping with Foundation Models", 20, color=MUTED)
box(s, 0.6, 1.0, 5.7, 1.9, "Щільна метрична 3D-карта з монокулярного RGB", 36, bold=True)
result = card(s, 0.6, 3.15, 5.6, 1.45, CARD)
box(s, 0.85, 3.3, 5.1, 1.2, "Побудовано відтворюваний offline pipeline: RGB → depth + SLAM → dense 3D map.",
    22, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, 0.6, 4.85, 5.7, 1.2, ["Learned metric depth (Depth Anything 3)", "+ класичний monocular SLAM (pySLAM)"], 20)
box(s, 0.6, 6.15, 5.7, 0.5, [[("Назар · Ярополк · Юліан", {"bold": True}),
                               (f"     develop @ {COMMIT}", {"color": MUTED, "size": 16})]], 20)
picture = s.shapes.add_picture(str(ASSETS / "final_dense_map.png"), Inches(6.55), Inches(1.0), width=Inches(6.3))
picture.crop_top, picture.crop_bottom = 50 / 720, 40 / 720  # drop the small burned-in caption lines
picture.height = int(picture.width * (720 - 90) / 1280)
tag = card(s, 6.55, 4.75, 3.2, 0.45, ORANGE, radius=0.3)
box(s, 6.55, 4.75, 3.2, 0.45, "Offline research prototype", 16, bold=True, color=BG,
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
box(s, 6.55, 5.35, 6.3, 1.2, ["Реальний результат: TUM RGB-D fr1/xyz, 200 кадрів,",
                              "22 keyframes, 72 887 вокселів по 2 см (run yulian_tum_e2e_f200_20261025_01)"],
    16, color=MUTED)
notes(s, 1, """ДЖЕРЕЛА: зображення — presentation/midterm/assets/final_dense_map.png, відрендерено scripts/render_midterm_demo.py з run yulian_tum_e2e_f200_20261025_01 (clean 4952e37, status ok); 22 keyframes і 72 887 вокселів — methods/…comparison.global/metrics.json цього run, results/summaries/wp3_wp4_yulian_tum_fr1_xyz_200.json, docs/ROBUST_FUSION.md §4.2. Replay карти побайтово збігається з записаним map.npz.""")

# ---------------------------------------------------------------- slide 2
s = slide_base("Архітектура і внесок команди", 2)
node(s, 0.6, 2.3, 1.75, 1.1, "RGB кадри", "monocular, 640×480")
node(s, 2.95, 1.3, 3.3, 1.2, "Metric depth", "DA3METRIC-LARGE (primary) · UniDepth V2 (порівняння)", "Назар")
node(s, 2.95, 3.2, 3.3, 1.2, "pySLAM monocular", "пози, keyframes, sparse map", "Ярополк")
node(s, 6.75, 2.3, 1.85, 1.1, "Alignment", "robust global scale", "Юліан")
node(s, 9.05, 2.3, 1.75, 1.1, "Dense fusion", "voxel 2 см", "Юліан")
node(s, 11.2, 2.3, 1.55, 1.1, "3D map", "метри", "Юліан")
node(s, 9.05, 3.8, 3.7, 0.95, "Evaluation з GT (TUM)", "лише оцінка: ATE, AbsRel, Chamfer, F‑score",
     line=MUTED, dashed=True)
arrow(s, 2.35, 2.65, 2.95, 1.9)
arrow(s, 2.35, 3.05, 2.95, 3.8)
arrow(s, 6.25, 1.9, 6.75, 2.65)
arrow(s, 6.25, 3.8, 6.75, 3.05)
arrow(s, 8.6, 2.85, 9.05, 2.85)
arrow(s, 10.8, 2.85, 11.2, 2.85)
arrow(s, 11.97, 3.4, 11.97, 3.8, dashed=True)
roles = [("Назар", "датасет і контракти, depth adapters, вибір моделі, інтеграція, бенчмарки"),
         ("Ярополк", "pySLAM, tracking, траєкторія, timing, SLAM-бенчмарк"),
         ("Юліан", "alignment, fusion, оцінка карти, robust-fusion дослідження")]
for i, (name, role) in enumerate(roles):
    x = 0.6 + i * 4.1
    card(s, x, 5.05, 3.85, 1.75, CARD, line=PEOPLE[name])
    box(s, x + 0.2, 5.15, 3.45, 0.45, name, 22, bold=True, color=PEOPLE[name])
    box(s, x + 0.2, 5.6, 3.45, 1.15, role, 20)
notes(s, 2, """ДЖЕРЕЛА: docs/ARCHITECTURE.md §1, §4 (Case A, scalar scale), §8; docs/EXPERIMENTS.md §3 (GT лише evaluation або labelled oracle B1); ролі — PLAN.md (WP1 Nazar, WP2 Yaropolk, WP3 Yulian), docs/TWO_PERSON_COMPLETION_PLAN.md §1.1, merged PR #1–#28. GT-free runtime перевіряється import-hygiene тестами (tests/test_import_hygiene_robust.py, tests/test_missing_gt_pose.py).""")

# ---------------------------------------------------------------- slide 3
s = slide_base("Depth і SLAM: якість моделей близька, latency — ні", 3)
box(s, 0.6, 1.1, 6.0, 0.4, "Development: TUM fr1/desk, 573 кадри", 18, color=MUTED)
data = CategoryChartData()
data.categories = ["AbsRel", "RMSE, м"]
data.add_series("DA3 (504)", (0.0819, 0.1626))
data.add_series("UniDepth V2", (0.0827, 0.1374))
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.45), Inches(6.0), Inches(3.55),
                           data).chart
style_chart(chart, "0.000", title="Похибка глибини (менше = краще)", legend=True, max_value=0.2, axis_format="0.00")
chart.plots[0].series[0].format.fill.solid()
chart.plots[0].series[0].format.fill.fore_color.rgb = ORANGE
chart.plots[0].series[1].format.fill.solid()
chart.plots[0].series[1].format.fill.fore_color.rgb = BLUE
chart.plots[0].gap_width = 60
data = CategoryChartData()
data.categories = ["бюджет 30 FPS", "UniDepth V2", "DA3 (504)", "pySLAM track"]
data.add_series("p95, ms", (33.3, 86.4, 97.9, 2312.5))
chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.9), Inches(1.45), Inches(5.9), Inches(3.55),
                           data).chart
style_chart(chart, "#,##0.0", title="p95 latency на кадр, ms", max_value=2800, axis_format="#,##0")
color_points(chart.plots[0].series[0], [MUTED, BLUE, ORANGE, GREY])
chart.plots[0].gap_width = 50
stats = [("430 / 573", "кадрів tracked (0.750)"), ("0.0195 м", "ATE після Sim(3), лише оцінка"),
         ("611 ms", "SLAM track, медіана на кадр")]
for i, (big, small) in enumerate(stats):
    x = 0.6 + i * 4.1
    card(s, x, 5.2, 3.85, 1.15, CARD)
    box(s, x + 0.2, 5.25, 3.5, 0.6, big, 28, bold=True)
    box(s, x + 0.2, 5.8, 3.5, 0.5, small, 18, color=MUTED)
box(s, 0.6, 6.45, 12.2, 0.5, "DA3 — primary за правилом D‑3 · SLAM домінує в latency → не real-time",
    20, bold=True, color=ORANGE)
notes(s, 3, """ДЖЕРЕЛА: depth — results/summaries/nazar_n05_model_only_0dc5f7f.csv (DA3 504: AbsRel 0.08189, RMSE 0.16258 м, p95 97.87 ms; UniDepth: AbsRel 0.08267, RMSE 0.13735 м, p95 86.42 ms; RTX 5070 Laptop, 640×480 input), docs/DEVELOPMENT_BENCHMARK.md §8; selection (DA3 AbsRel 0.1145 / RMSE 0.4211; UniDepth 0.0646 / 0.4499) — той самий CSV. B0 — results/summaries/nazar_n05_benchmark.json, run nazar_n05_dev573_b0eval_cd64251_01: tracking_success 0.75044 (430/573), final Sim(3) ATE RMSE 0.019530 м, Slam.track warm p50 611.38 / p95 2312.47 ms (Linux Docker, Python core); docs/DEVELOPMENT_BENCHMARK.md §4. D-3: docs/decisions/ADR001_geometry_model.md (N-04 gate fail → DA3 primary). Бюджет 33.3 ms = 1000/30 (D-7). Розбіжність: Y-05 Ярополка (results/summaries/yaropolk_y05_benchmark.csv, інша машина, 3 runs) дає track p95 1221–1339 ms і tracking 0.750–0.766; висновок «offline» той самий. Різниця між runs — також прояв недетермінізму pySLAM.""")

# ---------------------------------------------------------------- slide 4
s = slide_base("Демо: від RGB до щільної карти", 4)
video_w = 9.0
video_h = video_w * 720 / 1280
movie = s.shapes.add_movie(str(ASSETS / "mapping_demo.mp4"), Inches(0.6), Inches(1.25), Inches(video_w),
                           Inches(video_h), poster_frame_image=str(ASSETS / "mapping_demo_thumbnail.png"),
                           mime_type="video/mp4")
movie.name = "mapping_demo"
x = 9.95
box(s, x, 1.25, 2.85, 4.6, [
    [("Що на відео", {"bold": True, "color": ORANGE})],
    [("RGB → depth DA3 → траєкторія pySLAM → карта росте з keyframes", {"after": 10})],
    [("22 keyframes", {"bold": True})],
    [("72 887 вокселів", {"bold": True, "after": 10})],
    [("Rendered offline: швидкість відео ≠ швидкість обробки", {"color": ORANGE, "bold": True})],
], 20)
link = box(s, 0.6, 6.45, 9.0, 0.4, "", 14, color=MUTED)
paragraph = link.text_frame.paragraphs[0]
run = paragraph.add_run()
run.text = "Резерв: assets/mapping_demo.mp4 (31 с) · run yulian_tum_e2e_f200_20261025_01 (4952e37)"
run.font.size, run.font.name, run.font.color.rgb = Pt(14), FONT, MUTED
run.hyperlink.address = "assets/mapping_demo.mp4"
notes(s, 4, """ДЖЕРЕЛА: run yulian_tum_e2e_f200_20261025_01 (clean 4952e37, всі stages ok; DA3METRIC-LARGE process_res 384, bfloat16). Відео: scripts/render_midterm_demo.py, presentation/midterm/assets/demo_provenance.json (SHA-256 входів і виходів; replay_equals_recorded_map_npz = true). Tracking state не записаний: run передує Y-01 export; показано лише наявність online pose (197/200 кадрів). Траєкторія — trajectory_online.txt × scale 1.1920834 (alignment.json).""")

# ---------------------------------------------------------------- slide 5
s = slide_base("Карта: dense > sparse; per-frame anchoring відхилено", 5)
box(s, 0.6, 1.1, 12.2, 0.4, "Smoke: TUM fr1/xyz, 200 кадрів, 22 keyframes, DA3 (384) · Chamfer-L1, м — менше = краще",
    18, color=MUTED)
data = CategoryChartData()
data.categories = ["B0 sparse SLAM", "P1-global (наш)", "B1 oracle (GT depth)"]
data.add_series("Chamfer-L1, м", (0.1377, 0.0887, 0.0329))
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6.2), Inches(3.6),
                           data).chart
style_chart(chart, "0.0000", title="A. Основне порівняння карт", max_value=0.16, axis_format="0.00")
color_points(chart.plots[0].series[0], [GREY, ORANGE, ORACLE])
chart.plots[0].gap_width = 60
data = CategoryChartData()
data.categories = ["P1-global", "per-frame anchored"]
data.add_series("Chamfer-L1, м", (0.0887, 0.1051))
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.0), Inches(1.5), Inches(5.8), Inches(3.6),
                           data).chart
style_chart(chart, "0.0000", title="B. Per-frame anchoring (U-03)", max_value=0.16, axis_format="0.00")
color_points(chart.plots[0].series[0], [ORANGE, RED])
chart.plots[0].gap_width = 80
bullets(s, 0.6, 5.3, 6.1, 1.6, ["Dense краща за sparse (виграш у повноті)", "B1: ті самі пози і scale + GT depth → розрив = помилка depth"], 20)
bullets(s, 7.0, 5.3, 5.8, 1.6, ["F@5 см: 0.480 → 0.407", "Узгодженість зі SLAM ↑, точність карти ↓ → відхилено"], 20)
notes(s, 5, """ДЖЕРЕЛА: results/summaries/yulian_u03_robust_smoke200_cc7a4d9.csv (run yulian_u03_robust_smoke200_cc7a4d9_01, clean cc7a4d9): Chamfer-L1 B0-shared-scale 0.13774, P1-global control 0.08866, B1-oracle 0.03290, P1-anchored 0.10514; F@5cm 0.47997 → 0.40700; accuracy B0 0.0494 м проти P1-global 0.0924 м, completeness 0.2261 проти 0.0849 м, eval points 1 038 проти 58 569 (тому B0 «точний, але неповний»); P1-global control побайтово = записаному P1-global run yulian_tum_e2e_f200_20261025_01. Withheld residual 0.0388 → 0.0301 м — results/summaries/yulian_u02_anchor_gate_smoke200_2e6fad5.json; висновки — docs/ROBUST_FUSION.md §4.2, §6. Failure cases (кадри 80, 72, 84) — assets/failure_analysis.png, docs/ROBUST_FUSION.md §5. Smoke-спліт — DA3 process_res 384, не frozen 504.""")

# ---------------------------------------------------------------- slide 6
s = slide_base("Статус, обмеження і наступний крок", 6)
columns = [
    ("Зроблено", TEAL, ["End-to-end offline pipeline", "Порівняння depth‑моделей", "SLAM + dense fusion",
                        "Baselines B0/B1 і ablations", "Protocol freeze (IC3)",
                        "Manifests, run.json, tests, CI"]),
    ("Обмеження", RED, ["Offline, не 30 FPS", "pySLAM: недетермінізм, висока latency",
                        "Depth — найбільший розрив якості", "UniDepth обмежений правилом D‑3, зафіксованим заздалегідь",
                        "Перший held‑out запуск зупинився на preflight, без метрик"]),
    ("Далі", ORANGE, ["Preflight fix уже merged (PR #27, #29)",
                      "Один дозволений held‑out rerun із тим самим frozen config",
                      "Фінальний demo / integration audit",
                      "Новий протокол може переглянути UniDepth"]),
]
for i, (head, color, items) in enumerate(columns):
    x = 0.6 + i * 4.1
    card(s, x, 1.3, 3.85, 5.5, CARD, line=color)
    box(s, x + 0.25, 1.45, 3.4, 0.55, head, 24, bold=True, color=color)
    bullets(s, x + 0.25, 2.1, 3.4, 4.6, items, 20, gap=10)
notes(s, 6, """ДЖЕРЕЛА: docs/HELDOUT_BENCHMARK.md (run nazar_n06_heldout2488_da3_e561d8c_01: failed на WP2 preflight, 0 SLAM frames, 0 DA3 inference; results/summaries/nazar_n06_heldout2488_e561d8c.json); fix merged: PR #27 (c08833f) і PR #29 (242c0cb, partial GT coverage при 20 ms); N-06 rerun ще не запускався; docs/BENCHMARK_PROTOCOL.md (IC3 freeze, D-5 rerun); docs/DEVELOPMENT_BENCHMARK.md §5 (UniDepth P1 Chamfer 0.0447 vs DA3 0.1057 на dev; D-3 не переглядається); CI — .github/workflows/tests.yml. Окремо: Y-06 (PR #28) — B0-only SLAM held-out run Ярополка завершився; його числа на слайдах не показуємо, бо повного P1 held-out ще немає (див. README).""")

theme = prs.slide_master.part.part_related_by(RT.THEME)
theme._blob = re.sub(rb'(<a:(?:hlink|folHlink)><a:srgbClr val=")[0-9A-Fa-f]{6}', rb"\g<1>FFA826", theme.blob)
prs.save(OUT)
print(f"wrote {OUT.name}: {len(prs.slides)} slides")
